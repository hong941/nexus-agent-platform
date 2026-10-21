from __future__ import annotations

import hashlib
import json
import math
import re
import sqlite3
import threading
import uuid
from array import array
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .config import DATA_DIR, UPLOAD_DIR, settings


DATABASE_PATH = DATA_DIR / "knowledge.db"
LEGACY_STORE_PATH = DATA_DIR / "knowledge.json"


@dataclass(slots=True)
class Document:
    id: str
    title: str
    filename: str
    content: str
    size: int
    created_at: str


@dataclass(slots=True)
class Chunk:
    id: str
    document_id: str
    title: str
    content: str
    chunk_index: int
    vector: list[float]


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _normalize(text: str) -> str:
    without_headings = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.MULTILINE)
    return re.sub(r"\s+", " ", without_headings).strip()


class HashEmbedding:
    """Deterministic local embedding used when bge-m3 is unavailable."""

    dimensions = 384

    def encode(self, text: str) -> list[float]:
        vector = [0.0] * self.dimensions
        compact = re.sub(r"\s+", "", text.lower())
        features = [compact[index : index + 2] for index in range(max(1, len(compact) - 1))]
        features.extend(re.findall(r"[a-z0-9_]+", text.lower()))
        for feature in features:
            digest = hashlib.blake2b(feature.encode("utf-8"), digest_size=8).digest()
            index = int.from_bytes(digest[:4], "little") % self.dimensions
            sign = 1.0 if digest[4] % 2 == 0 else -1.0
            vector[index] += sign
        norm = math.sqrt(sum(value * value for value in vector)) or 1.0
        return [value / norm for value in vector]


class BGEEmbedding:
    def __init__(self, model_name: str) -> None:
        from sentence_transformers import SentenceTransformer  # type: ignore[import-not-found]

        self.model = SentenceTransformer(model_name)

    def encode(self, text: str) -> list[float]:
        encoded = self.model.encode(text, normalize_embeddings=True)
        return encoded.tolist()


def _chunk_text(text: str, size: int = 520, overlap: int = 90) -> list[str]:
    clean = text.replace("\r\n", "\n").strip()
    if not clean:
        return []
    paragraphs = [segment.strip() for segment in re.split(r"\n{2,}", clean) if segment.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(current) + len(paragraph) + 2 <= size:
            current = f"{current}\n\n{paragraph}".strip()
            continue
        if current:
            chunks.append(current)
        if len(paragraph) <= size:
            current = paragraph
            continue
        start = 0
        while start < len(paragraph):
            piece = paragraph[start : start + size]
            chunks.append(piece.strip())
            start += size - overlap
        current = ""
    if current:
        chunks.append(current)
    return [chunk for chunk in chunks if chunk]


class KnowledgeBase:
    def __init__(self) -> None:
        self._lock = threading.RLock()
        self.documents: dict[str, Document] = {}
        self.chunks: dict[str, Chunk] = {}
        self.embedding_name = "local-hash-384"
        self.embedding: HashEmbedding | BGEEmbedding = HashEmbedding()
        self.chroma_client: Any = None
        self.collection: Any = None
        self.backend = "sqlite"
        self._connection = self._open_database()

        if settings.embedding_model:
            try:
                self.embedding = BGEEmbedding(settings.embedding_model)
                self.embedding_name = settings.embedding_model
            except Exception:
                self.embedding = HashEmbedding()

        self._load()
        if not self.documents:
            self._migrate_legacy_json()
            self._load()

        if settings.use_chroma:
            try:
                import chromadb

                self.chroma_client = chromadb.PersistentClient(path=str(DATA_DIR / "chroma"))
                self.collection = self.chroma_client.get_or_create_collection(
                    name="nexus_knowledge",
                    metadata={"hnsw:space": "cosine"},
                )
                self.backend = "chromadb"
                self._sync_chroma()
            except Exception:
                self.collection = None
                self.backend = "sqlite"

        if not self.documents:
            self._seed()

    @staticmethod
    def _open_database() -> sqlite3.Connection:
        connection = sqlite3.connect(DATABASE_PATH, check_same_thread=False)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA journal_mode=WAL")
        connection.execute("PRAGMA foreign_keys=ON")
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS documents (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                filename TEXT NOT NULL,
                content TEXT NOT NULL,
                size INTEGER NOT NULL,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS chunks (
                id TEXT PRIMARY KEY,
                document_id TEXT NOT NULL,
                title TEXT NOT NULL,
                content TEXT NOT NULL,
                chunk_index INTEGER NOT NULL,
                vector BLOB NOT NULL,
                FOREIGN KEY (document_id) REFERENCES documents(id) ON DELETE CASCADE
            );

            CREATE INDEX IF NOT EXISTS idx_chunks_document
            ON chunks(document_id);
            """
        )
        connection.commit()
        return connection

    def _load(self) -> None:
        with self._lock:
            self.documents.clear()
            self.chunks.clear()
            for row in self._connection.execute("SELECT * FROM documents"):
                document = Document(
                    id=row["id"],
                    title=row["title"],
                    filename=row["filename"],
                    content=row["content"],
                    size=row["size"],
                    created_at=row["created_at"],
                )
                self.documents[document.id] = document
            for row in self._connection.execute("SELECT * FROM chunks"):
                chunk = Chunk(
                    id=row["id"],
                    document_id=row["document_id"],
                    title=row["title"],
                    content=row["content"],
                    chunk_index=row["chunk_index"],
                    vector=self._blob_to_vector(row["vector"]),
                )
                self.chunks[chunk.id] = chunk

    @staticmethod
    def _vector_to_blob(vector: list[float]) -> bytes:
        return array("f", vector).tobytes()

    @staticmethod
    def _blob_to_vector(payload: bytes) -> list[float]:
        values = array("f")
        values.frombytes(payload)
        return values.tolist()

    def _migrate_legacy_json(self) -> None:
        if not LEGACY_STORE_PATH.is_file():
            return
        try:
            payload = json.loads(LEGACY_STORE_PATH.read_text(encoding="utf-8"))
            documents = [Document(**item) for item in payload.get("documents", [])]
            chunks = [Chunk(**item) for item in payload.get("chunks", [])]
            if not documents:
                return
            with self._connection:
                self._connection.executemany(
                    """
                    INSERT OR IGNORE INTO documents
                    (id, title, filename, content, size, created_at)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            item.id,
                            item.title,
                            item.filename,
                            item.content,
                            item.size,
                            item.created_at,
                        )
                        for item in documents
                    ],
                )
                self._connection.executemany(
                    """
                    INSERT OR IGNORE INTO chunks
                    (id, document_id, title, content, chunk_index, vector)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            item.id,
                            item.document_id,
                            item.title,
                            item.content,
                            item.chunk_index,
                            self._vector_to_blob(item.vector),
                        )
                        for item in chunks
                    ],
                )
            migrated = LEGACY_STORE_PATH.with_suffix(".json.migrated")
            if not migrated.exists():
                LEGACY_STORE_PATH.replace(migrated)
        except (OSError, ValueError, TypeError, sqlite3.Error):
            # Keep the legacy file untouched when migration cannot complete.
            return

    def _sync_chroma(self) -> None:
        if self.collection is None or not self.chunks:
            return
        chunk_list = list(self.chunks.values())
        self.collection.upsert(
            ids=[chunk.id for chunk in chunk_list],
            documents=[chunk.content for chunk in chunk_list],
            embeddings=[chunk.vector for chunk in chunk_list],
            metadatas=[
                {
                    "document_id": chunk.document_id,
                    "title": chunk.title,
                    "chunk_index": chunk.chunk_index,
                }
                for chunk in chunk_list
            ],
        )

    def _seed(self) -> None:
        seeds = [
            (
                "Nexus Agent 产品说明",
                "nexus-product.md",
                """# Nexus Agent 多 Agent 智能对话平台

Nexus Agent 使用 Supervisor 编排数学、代码和知识三个专业 Agent。Supervisor 先分析问题意图，再将任务分发给最合适的 Agent；复杂问题可以顺序调用多个 Agent，最后由 Synthesis Agent 合并结果。

平台支持 SSE 流式输出。前端能实时看到路由决策、工具调用、检索来源和最终回答，而不是只等待一个长请求结束。对话记录保存在浏览器中，因此刷新页面后仍可继续最近的会话。

知识库支持 Markdown、TXT、PDF 和 Word 文档。上传的文档会被切分、向量化并写入检索库。回答时会返回命中的文档标题、片段和相关度分数。""",
            ),
            (
                "RAG 工程规范",
                "rag-guide.md",
                """# RAG 知识库规范

文档切块采用约 520 字符、90 字符重叠的策略，兼顾上下文完整性与召回精度。默认演示环境使用确定性向量，便于离线验证完整链路；设置 USE_CHROMA=true 并安装可选依赖后，会使用 ChromaDB 持久化向量。

生产环境可以设置 EMBEDDING_MODEL=BAAI/bge-m3 启用 bge-m3 嵌入。检索结果需要在回答中展示来源引用，并避免把低相关度片段当成确定事实。

知识类 Agent 的提示词要求：优先依据检索片段回答；资料不足时明确说明；不得捏造来源。""",
            ),
            (
                "平台使用技巧",
                "usage-tips.md",
                """# 使用技巧

需要计算时可以直接输入数学表达式，例如“计算 (128 * 7 + 36) / 4”。数学 Agent 会调用受限计算器，而不是让语言模型猜测结果。

需要写代码时可以说明语言、输入输出和边界条件。代码 Agent 会给出可运行示例，并调用代码检查器检查明显问题。

询问平台功能、RAG 原理或已上传资料时，问题会自动路由到知识 Agent，并在回答下方展示命中的知识来源。""",
            ),
        ]
        for title, filename, content in seeds:
            self.add_document(title=title, filename=filename, content=content)

    def add_document(self, title: str, filename: str, content: str) -> Document:
        clean = content.strip()
        if not clean:
            raise ValueError("文档内容为空")
        document = Document(
            id=uuid.uuid4().hex,
            title=title.strip() or filename,
            filename=filename,
            content=clean,
            size=len(clean.encode("utf-8")),
            created_at=_now_iso(),
        )
        new_chunks: list[Chunk] = []
        for index, text in enumerate(_chunk_text(clean)):
            new_chunks.append(
                Chunk(
                    id=f"{document.id}:{index}",
                    document_id=document.id,
                    title=document.title,
                    content=text,
                    chunk_index=index,
                    vector=self.embedding.encode(f"{document.title}\n{text}"),
                )
            )
        with self._lock:
            with self._connection:
                self._connection.execute(
                    """
                    INSERT INTO documents
                    (id, title, filename, content, size, created_at)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (
                        document.id,
                        document.title,
                        document.filename,
                        document.content,
                        document.size,
                        document.created_at,
                    ),
                )
                self._connection.executemany(
                    """
                    INSERT INTO chunks
                    (id, document_id, title, content, chunk_index, vector)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            chunk.id,
                            chunk.document_id,
                            chunk.title,
                            chunk.content,
                            chunk.chunk_index,
                            self._vector_to_blob(chunk.vector),
                        )
                        for chunk in new_chunks
                    ],
                )
            self.documents[document.id] = document
            for chunk in new_chunks:
                self.chunks[chunk.id] = chunk
            if self.collection is not None and new_chunks:
                try:
                    self.collection.upsert(
                        ids=[chunk.id for chunk in new_chunks],
                        documents=[chunk.content for chunk in new_chunks],
                        embeddings=[chunk.vector for chunk in new_chunks],
                        metadatas=[
                            {
                                "document_id": chunk.document_id,
                                "title": chunk.title,
                                "chunk_index": chunk.chunk_index,
                            }
                            for chunk in new_chunks
                        ],
                    )
                except Exception:
                    self.collection = None
                    self.backend = "sqlite"
        return document

    def delete_document(self, document_id: str) -> bool:
        with self._lock:
            if document_id not in self.documents:
                return False
            document = self.documents[document_id]
            chunk_ids = [
                chunk_id
                for chunk_id, chunk in self.chunks.items()
                if chunk.document_id == document_id
            ]
            with self._connection:
                self._connection.execute(
                    "DELETE FROM documents WHERE id = ?", (document_id,)
                )
            for chunk_id in chunk_ids:
                self.chunks.pop(chunk_id, None)
            self.documents.pop(document_id, None)
            if self.collection is not None and chunk_ids:
                try:
                    self.collection.delete(ids=chunk_ids)
                except Exception:
                    self.collection = None
                    self.backend = "sqlite"
            upload = UPLOAD_DIR / f"{document_id}_{document.filename}"
            if upload.is_file():
                upload.unlink()
            return True

    def search(self, query: str, top_k: int = 4) -> list[dict[str, Any]]:
        if not query.strip():
            return []
        with self._lock:
            if self.collection is not None and self.chunks:
                try:
                    result = self.collection.query(
                        query_embeddings=[self.embedding.encode(query)],
                        n_results=min(top_k, len(self.chunks)),
                    )
                    ids = (result.get("ids") or [[]])[0]
                    distances = (result.get("distances") or [[]])[0]
                    matches = []
                    for index, chunk_id in enumerate(ids):
                        chunk = self.chunks.get(chunk_id)
                        if not chunk:
                            continue
                        distance = (
                            float(distances[index]) if index < len(distances) else 1.0
                        )
                        matches.append(
                            self._match_dict(chunk, 1.0 / (1.0 + distance))
                        )
                    return matches
                except Exception:
                    self.collection = None
                    self.backend = "sqlite"

            query_vector = self.embedding.encode(query)
            scored: list[tuple[float, Chunk]] = []
            terms = set(re.findall(r"[\w\u4e00-\u9fff]{2,}", query.lower()))
            for chunk in self.chunks.values():
                cosine = sum(a * b for a, b in zip(query_vector, chunk.vector))
                haystack = f"{chunk.title} {chunk.content}".lower()
                lexical = sum(1 for term in terms if term in haystack) / max(1, len(terms))
                score = 0.72 * max(0.0, cosine) + 0.28 * lexical
                scored.append((score, chunk))
            scored.sort(key=lambda item: item[0], reverse=True)
            # Normalize small local-vector scores into a readable relevance range.
            top = scored[0][0] if scored else 0.0
            matches = []
            for raw_score, chunk in scored[:top_k]:
                normalized = raw_score / top if top else 0.0
                if normalized >= 0.18:
                    matches.append(self._match_dict(chunk, normalized))
            return matches

    @staticmethod
    def _match_dict(chunk: Chunk, score: float) -> dict[str, Any]:
        return {
            "id": chunk.id,
            "document_id": chunk.document_id,
            "title": chunk.title,
            "chunk_index": chunk.chunk_index,
            "excerpt": _normalize(chunk.content)[:260],
            "score": round(max(0.0, min(1.0, score)), 3),
        }

    def list_documents(self) -> list[dict[str, Any]]:
        with self._lock:
            rows = self._connection.execute(
                """
                SELECT d.*, COUNT(c.id) AS chunk_count
                FROM documents d
                LEFT JOIN chunks c ON c.document_id = d.id
                GROUP BY d.id
                ORDER BY d.created_at DESC
                """
            ).fetchall()
            return [
                {
                    "id": row["id"],
                    "title": row["title"],
                    "filename": row["filename"],
                    "chunks": row["chunk_count"],
                    "size": row["size"],
                    "created_at": row["created_at"],
                }
                for row in rows
            ]

    def stats(self) -> dict[str, Any]:
        document_count = self._connection.execute(
            "SELECT COUNT(*) FROM documents"
        ).fetchone()[0]
        chunk_count = self._connection.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
        return {
            "documents": document_count,
            "chunks": chunk_count,
            "backend": self.backend,
            "database": DATABASE_PATH.name,
            "embedding": self.embedding_name,
        }


def extract_text(filename: str, payload: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".csv"}:
        for encoding in ("utf-8", "gb18030", "latin-1"):
            try:
                return payload.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("无法识别文本编码")
    if suffix == ".pdf":
        from io import BytesIO
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(payload))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        from io import BytesIO
        from docx import Document as DocxDocument

        document = DocxDocument(BytesIO(payload))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    raise ValueError("仅支持 Markdown、TXT、PDF、DOCX 和 CSV")


knowledge_base = KnowledgeBase()

